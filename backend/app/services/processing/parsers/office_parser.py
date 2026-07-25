"""
Industrial Office Parser for DOCX, XLSX, CSV, and PPTX records.
Extracts work orders, SOP operating manuals, maintenance logs, and parameter spreadsheets.
"""
import os
import re
import logging
from typing import List, Dict, Any, Tuple

from .base import BaseParser, ParsedDocument, ParsedChunk, ParsedTable

logger = logging.getLogger("industrybrain.parsers.office")

EQUIPMENT_TAG_REGEX = re.compile(r"\b([A-Z]{1,4}-[0-9]{2,4}[A-Z]?)\b")
REGULATORY_REGEX = re.compile(r"\b(OISD-[0-9]{3}|PESO|Factory Act|ISO\s?[0-9]{4,5}|OSHA)\b", re.IGNORECASE)


class OfficeParser(BaseParser):
    def parse(self, file_path: str, filename: str) -> ParsedDocument:
        ext = os.path.splitext(filename)[1].lower()
        
        if ext == ".docx":
            return self._parse_docx(file_path, filename)
        elif ext in [".xlsx", ".xls", ".csv"]:
            return self._parse_spreadsheet(file_path, filename, ext)
        elif ext == ".pptx":
            return self._parse_pptx(file_path, filename)
        else:
            raise ValueError(f"Unsupported office extension: {ext}")

    def _parse_docx(self, file_path: str, filename: str) -> ParsedDocument:
        logger.info(f"[OfficeParser] Extracting Word document '{filename}'")
        import docx

        doc = docx.Document(file_path)
        chunks = []
        tables = []
        raw_text_parts = []
        equipment_tags = set()
        regulatory_refs = set()
        chunk_idx = 0

        # Paragraphs & Headings
        current_heading = "General"
        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue

            if p.style.name.startswith("Heading"):
                current_heading = text

            raw_text_parts.append(text)
            
            for tag in EQUIPMENT_TAG_REGEX.findall(text):
                equipment_tags.add(tag)
            for reg in REGULATORY_REGEX.findall(text):
                regulatory_refs.add(reg.upper())

            if len(text) > 30:
                chunks.append(ParsedChunk(
                    content=text,
                    chunk_index=chunk_idx,
                    section_heading=current_heading,
                    metadata={"source": filename, "style": p.style.name}
                ))
                chunk_idx += 1

        # Word Tables
        for t_idx, table in enumerate(doc.tables):
            matrix = []
            for row in table.rows:
                row_data = [cell.text.strip() for cell in row.cells]
                matrix.append(row_data)

            if len(matrix) >= 2:
                tables.append(ParsedTable(
                    title=f"Table {t_idx+1}",
                    headers=matrix[0],
                    rows=matrix[1:],
                    page_number=1
                ))

        return ParsedDocument(
            filename=filename,
            file_type="docx",
            raw_text="\n\n".join(raw_text_parts),
            chunks=chunks,
            tables=tables,
            equipment_tags=list(equipment_tags),
            regulatory_references=list(regulatory_refs)
        )

    def _parse_spreadsheet(self, file_path: str, filename: str, ext: str) -> ParsedDocument:
        logger.info(f"[OfficeParser] Extracting Excel/CSV spreadsheet '{filename}'")
        chunks = []
        tables = []
        raw_text_parts = []
        equipment_tags = set()
        chunk_idx = 0

        if ext == ".csv":
            import csv
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                reader = csv.reader(f)
                rows = list(reader)
                if rows:
                    tables.append(ParsedTable(
                        title="CSV Log Sheet",
                        headers=rows[0],
                        rows=rows[1:]
                    ))
                    for r in rows:
                        row_str = " | ".join(r)
                        raw_text_parts.append(row_str)
                        for tag in EQUIPMENT_TAG_REGEX.findall(row_str):
                            equipment_tags.add(tag)
        else:
            import openpyxl
            wb = openpyxl.load_workbook(file_path, data_only=True)
            for sheetname in wb.sheetnames:
                sheet = wb[sheetname]
                sheet_rows = []
                for row in sheet.iter_rows(values_only=True):
                    row_vals = [str(val or "").strip() for val in row if val is not None]
                    if any(row_vals):
                        sheet_rows.append(row_vals)
                        row_str = " | ".join(row_vals)
                        raw_text_parts.append(row_str)
                        for tag in EQUIPMENT_TAG_REGEX.findall(row_str):
                            equipment_tags.add(tag)

                if len(sheet_rows) >= 2:
                    tables.append(ParsedTable(
                        title=f"Sheet: {sheetname}",
                        headers=sheet_rows[0],
                        rows=sheet_rows[1:]
                    ))
            wb.close()

        # Chunk spreadsheet lines
        full_text = "\n".join(raw_text_parts)
        for i in range(0, len(raw_text_parts), 15):
            batch = "\n".join(raw_text_parts[i:i+15])
            if batch.strip():
                chunks.append(ParsedChunk(
                    content=batch,
                    chunk_index=chunk_idx,
                    metadata={"source": filename}
                ))
                chunk_idx += 1

        return ParsedDocument(
            filename=filename,
            file_type="xlsx" if ext != ".csv" else "csv",
            raw_text=full_text,
            chunks=chunks,
            tables=tables,
            equipment_tags=list(equipment_tags)
        )

    def _parse_pptx(self, file_path: str, filename: str) -> ParsedDocument:
        logger.info(f"[OfficeParser] Extracting PowerPoint slides '{filename}'")
        import pptx

        prs = pptx.Presentation(file_path)
        raw_text_parts = []
        chunks = []
        equipment_tags = set()
        chunk_idx = 0

        for slide_idx, slide in enumerate(prs.slides):
            slide_num = slide_idx + 1
            slide_texts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    txt = shape.text.strip()
                    slide_texts.append(txt)
                    for tag in EQUIPMENT_TAG_REGEX.findall(txt):
                        equipment_tags.add(tag)

            slide_body = "\n".join(slide_texts)
            if slide_body:
                raw_text_parts.append(f"--- Slide {slide_num} ---\n" + slide_body)
                chunks.append(ParsedChunk(
                    content=slide_body,
                    chunk_index=chunk_idx,
                    page_number=slide_num,
                    section_heading=f"Slide {slide_num}",
                    metadata={"source": filename, "slide": slide_num}
                ))
                chunk_idx += 1

        return ParsedDocument(
            filename=filename,
            file_type="pptx",
            raw_text="\n\n".join(raw_text_parts),
            chunks=chunks,
            equipment_tags=list(equipment_tags)
        )
